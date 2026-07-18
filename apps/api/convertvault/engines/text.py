import re
import shutil
import subprocess
import tempfile
import unicodedata
import zipfile
from pathlib import Path
from xml.etree import ElementTree

import fitz
from bs4 import BeautifulSoup
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook
from PIL import Image, ImageOps, ImageSequence
from pptx import Presentation
from striprtf.striprtf import rtf_to_text

from ..config import settings
from .base import Capability, ConversionContext


class TextExtractionEngine:
    engine_id = "local-text-extractor"
    display_name = "Local structured text extraction"
    document_sources = ["docx", "odt", "rtf", "txt", "md", "markdown", "html", "htm", "csv", "tsv", "xlsx", "pptx", "epub"]
    image_sources = ["jpg", "jpeg", "png", "webp", "gif", "tiff", "tif", "bmp", "heic", "heif"]

    def available(self) -> bool:
        return True

    def version(self) -> str:
        return "1"

    def capabilities(self) -> list[Capability]:
        options = {
            "normalize_whitespace": {"type": "boolean", "default": True},
            "reflow_paragraphs": {"type": "boolean", "default": True},
            "ocr_language": {"type": "string", "default": "eng"},
        }
        result = [
            Capability("text.extract", self.document_sources, ["txt", "md"], self.engine_id,
                       options=options, limitations=["Complex visual layout is converted into reading order; source wording is not rewritten."]),
            Capability("text.extract", ["pdf"], ["txt", "md"], self.engine_id, approximate=True,
                       options=options, limitations=["Selectable text is preserved; pages without sufficient text use OCR when Tesseract is installed."]),
        ]
        if shutil.which("tesseract"):
            result.append(Capability("text.extract", self.image_sources, ["txt", "md"], self.engine_id,
                                     approximate=True, options=options,
                                     limitations=["OCR cannot guarantee 100% accuracy; results depend on image quality, language data, and typography."]))
        return result

    def convert(self, context: ConversionContext) -> list[Path]:
        extension = context.source_extension.lower()
        extractor = {
            "pdf": self._pdf,
            "docx": self._docx,
            "odt": self._odt,
            "rtf": self._rtf,
            "txt": self._plain,
            "md": self._plain,
            "markdown": self._plain,
            "html": self._html,
            "htm": self._html,
            "csv": self._delimited,
            "tsv": self._delimited,
            "xlsx": self._xlsx,
            "pptx": self._pptx,
            "epub": self._epub,
        }.get(extension)
        if extractor:
            text = extractor(context)
        elif extension in self.image_sources:
            text = self._image(context)
        else:
            raise ValueError(f"Text extraction is not supported for .{extension}")
        if context.options.get("normalize_whitespace", True):
            text = clean_text(text)
        if context.options.get("reflow_paragraphs", True):
            text = reflow_text(text)
        if context.target_extension == "txt":
            text = re.sub(r"^#{1,6} ", "", text, flags=re.MULTILINE)
        if not text.strip():
            raise ValueError("No readable text was found in this file")
        context.destination.write_text(text.strip() + "\n", encoding="utf-8")
        return [context.destination]

    def _plain(self, context: ConversionContext) -> str:
        return context.source.read_text(encoding="utf-8", errors="replace")

    def _html(self, context: ConversionContext) -> str:
        soup = BeautifulSoup(context.source.read_text(encoding="utf-8", errors="replace"), "html.parser")
        for element in soup(["script", "style", "noscript"]): element.decompose()
        return soup.get_text("\n")

    def _docx(self, context: ConversionContext) -> str:
        document = Document(context.source); parts: list[str] = []
        for block in document.iter_inner_content():
            if isinstance(block, Paragraph):
                text = block.text.strip()
                if text:
                    level = re.search(r"Heading (\d+)", block.style.name if block.style else "")
                    parts.append(("#" * int(level.group(1)) + " " if level else "") + text)
            elif isinstance(block, Table):
                parts.extend("\t".join(cell.text.strip() for cell in row.cells) for row in block.rows)
        for section in document.sections:
            header, footer = section.header.paragraphs, section.footer.paragraphs
            if any(p.text.strip() for p in header): parts.append("\n[Header]\n" + "\n".join(p.text for p in header if p.text.strip()))
            if any(p.text.strip() for p in footer): parts.append("\n[Footer]\n" + "\n".join(p.text for p in footer if p.text.strip()))
        return "\n\n".join(parts)

    def _odt(self, context: ConversionContext) -> str:
        with zipfile.ZipFile(context.source) as archive: root = ElementTree.fromstring(archive.read("content.xml"))
        parts: list[str] = []
        for element in root.iter():
            local = element.tag.rsplit("}", 1)[-1]
            if local in {"p", "h", "table-cell"}:
                value = "".join(element.itertext()).strip()
                if value: parts.append(value)
        return "\n".join(parts)

    def _rtf(self, context: ConversionContext) -> str:
        return rtf_to_text(context.source.read_text(encoding="utf-8", errors="replace"))

    def _delimited(self, context: ConversionContext) -> str:
        delimiter = "\t" if context.source_extension == "tsv" else ","
        import csv
        with context.source.open(encoding="utf-8-sig", errors="replace", newline="") as stream:
            return "\n".join("\t".join(row) for row in csv.reader(stream, delimiter=delimiter))

    def _xlsx(self, context: ConversionContext) -> str:
        workbook = load_workbook(context.source, read_only=True, data_only=True); parts: list[str] = []
        try:
            for sheet in workbook.worksheets:
                parts.append(f"# {sheet.title}")
                for row in sheet.iter_rows(values_only=True):
                    values = ["" if value is None else str(value) for value in row]
                    if any(values): parts.append("\t".join(values).rstrip())
        finally: workbook.close()
        return "\n".join(parts)

    def _pptx(self, context: ConversionContext) -> str:
        presentation = Presentation(context.source); parts: list[str] = []
        for index, slide in enumerate(presentation.slides, 1):
            parts.append(f"# Slide {index}")
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip(): parts.append(shape.text.strip())
                if getattr(shape, "has_table", False):
                    parts.extend("\t".join(cell.text.strip() for cell in row.cells) for row in shape.table.rows)
            if slide.has_notes_slide:
                notes = [shape.text.strip() for shape in slide.notes_slide.notes_text_frame.paragraphs if shape.text.strip()]
                if notes: parts.append("[Speaker notes]\n" + "\n".join(notes))
        return "\n\n".join(parts)

    def _epub(self, context: ConversionContext) -> str:
        parts: list[str] = []
        with zipfile.ZipFile(context.source) as archive:
            for name in sorted(archive.namelist()):
                if name.lower().endswith((".xhtml", ".html", ".htm")):
                    soup = BeautifulSoup(archive.read(name), "html.parser")
                    for element in soup(["script", "style"]): element.decompose()
                    value = soup.get_text("\n").strip()
                    if value: parts.append(value)
        return "\n\n".join(parts)

    def _pdf(self, context: ConversionContext) -> str:
        pages: list[str] = []
        with fitz.open(context.source) as document:
            if document.needs_pass: raise ValueError("The PDF is password-protected; provide an unlocked copy for extraction")
            for index, page in enumerate(document):
                blocks = sorted(page.get_text("blocks"), key=lambda block: (round(block[1] / 8), block[0]))
                selectable = "\n".join(block[4].strip() for block in blocks if block[4].strip())
                if len(selectable) >= 20:
                    pages.append(f"# Page {index + 1}\n{selectable}")
                elif shutil.which("tesseract"):
                    pixmap = page.get_pixmap(matrix=fitz.Matrix(300 / 72, 300 / 72), alpha=False)
                    pages.append(f"# Page {index + 1}\n{self._ocr_bytes(pixmap.tobytes('png'), context)}")
                else:
                    pages.append(f"# Page {index + 1}\n[No selectable text; OCR is unavailable]")
        return "\n\n".join(pages)

    def _image(self, context: ConversionContext) -> str:
        if not shutil.which("tesseract"): raise RuntimeError("Image OCR is unavailable because Tesseract is not installed")
        pages: list[str] = []
        with Image.open(context.source) as opened:
            for index, frame in enumerate(ImageSequence.Iterator(opened), 1):
                image = ImageOps.exif_transpose(frame).convert("RGB"); buffer = tempfile.SpooledTemporaryFile(max_size=32 * 1024 * 1024)
                image.save(buffer, "PNG"); buffer.seek(0); pages.append(f"# Image {index}\n{self._ocr_bytes(buffer.read(), context)}"); buffer.close()
        return "\n\n".join(pages)

    def _ocr_bytes(self, data: bytes, context: ConversionContext) -> str:
        language = str(context.options.get("ocr_language", "eng"))
        if not re.fullmatch(r"[A-Za-z0-9_+.-]{2,80}", language): raise ValueError("OCR language code is invalid")
        with tempfile.NamedTemporaryFile(suffix=".png") as image:
            image.write(data); image.flush()
            result = subprocess.run(["tesseract", image.name, "stdout", "-l", language, "--psm", "3"],
                                    capture_output=True, text=True, timeout=settings.job_timeout_pdf, check=False)
        if result.returncode != 0: raise RuntimeError(f"OCR failed: {(result.stderr or 'unknown error').strip()[-300:]}")
        return result.stdout


def clean_text(value: str) -> str:
    value = unicodedata.normalize("NFKC", value).replace("\u00ad", "")
    value = re.sub(r"(?<=\w)-\n(?=\w)", "", value)
    lines = ["\t".join(re.sub(r" +", " ", cell.strip()) for cell in line.split("\t"))
             for line in value.replace("\r\n", "\n").replace("\r", "\n").split("\n")]
    output: list[str] = []
    for line in lines:
        if line or not output or output[-1]: output.append(line)
    return "\n".join(output).strip()


def reflow_text(value: str) -> str:
    """Repair prose line wrapping without changing word order or punctuation."""
    output: list[str] = []
    structural = re.compile(r"^(#{1,6}\s|[-*•]\s|\d+[.)]\s|\[.+\]$)")
    for line in value.splitlines():
        if not line:
            if output and output[-1]: output.append("")
            continue
        previous = output[-1] if output else ""
        can_join = bool(previous and "\t" not in previous and "\t" not in line
                        and not structural.match(previous) and not structural.match(line)
                        and not re.search(r"[.!?;:]$", previous))
        if can_join:
            output[-1] = f"{previous} {line}"
        else:
            output.append(line)
    return "\n".join(output).strip()
