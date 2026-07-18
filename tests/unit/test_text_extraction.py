from docx import Document
from openpyxl import Workbook
from PIL import Image
from pptx import Presentation
import fitz

from convertvault.engines.base import ConversionContext
from convertvault.engines.text import TextExtractionEngine, clean_text, reflow_text


def extract(engine, source, target, extension, options=None):
    engine.convert(ConversionContext(source, target, extension, target.suffix.lstrip("."), options or {}))
    return target.read_text(encoding="utf-8")


def test_clean_text_preserves_wording_and_repairs_line_hyphenation():
    assert clean_text("A  carefully hyphen-\nated\n\n\nphrase") == "A carefully hyphenated\n\nphrase"
    assert reflow_text("A sentence that was\nwrapped across lines.\n\n# Heading\nNext thought.") == "A sentence that was wrapped across lines.\n\n# Heading\nNext thought."


def test_docx_paragraphs_headings_and_tables_are_extracted(tmp_path):
    source, target = tmp_path / "document.docx", tmp_path / "document.md"
    document = Document(); document.add_heading("Quarterly report", level=1); document.add_paragraph("Revenue increased.")
    table = document.add_table(rows=2, cols=2); table.cell(0, 0).text = "Region"; table.cell(0, 1).text = "Total"; table.cell(1, 0).text = "North"; table.cell(1, 1).text = "42"
    document.save(source)
    text = extract(TextExtractionEngine(), source, target, "docx")
    assert "# Quarterly report" in text and "Region\tTotal" in text and "North\t42" in text


def test_pdf_html_spreadsheet_and_presentation_extraction(tmp_path):
    engine = TextExtractionEngine()
    pdf = tmp_path / "source.pdf"; document = fitz.open(); page = document.new_page(); page.insert_text((72,72), "Selectable PDF sentence."); document.save(pdf)
    assert "Selectable PDF sentence." in extract(engine, pdf, tmp_path / "pdf.txt", "pdf")
    html = tmp_path / "source.html"; html.write_text("<h1>Heading</h1><p>Proper sentence.</p><script>hidden()</script>", encoding="utf-8")
    html_text = extract(engine, html, tmp_path / "html.txt", "html")
    assert "Proper sentence." in html_text and "hidden" not in html_text
    workbook = Workbook(); sheet = workbook.active; sheet.title = "Data"; sheet.append(["Name", "Value"]); sheet.append(["Alpha", 7]); xlsx = tmp_path / "sheet.xlsx"; workbook.save(xlsx)
    sheet_text = extract(engine, xlsx, tmp_path / "sheet.md", "xlsx")
    assert "# Data" in sheet_text and "Alpha\t7" in sheet_text
    presentation = Presentation(); slide = presentation.slides.add_slide(presentation.slide_layouts[1]); slide.shapes.title.text = "Slide title"; slide.placeholders[1].text = "Complete thought."
    pptx = tmp_path / "slides.pptx"; presentation.save(pptx)
    slide_text = extract(engine, pptx, tmp_path / "slides.txt", "pptx")
    assert "Slide title" in slide_text and "Complete thought." in slide_text


def test_image_ocr_path_supports_multi_frame_reading_order(monkeypatch, tmp_path):
    source, target = tmp_path / "scan.tiff", tmp_path / "scan.txt"
    first, second = Image.new("RGB", (20,20), "white"), Image.new("RGB", (20,20), "gray")
    first.save(source, save_all=True, append_images=[second])
    engine = TextExtractionEngine(); calls = []
    monkeypatch.setattr("shutil.which", lambda name: "/usr/bin/tesseract")
    monkeypatch.setattr(engine, "_ocr_bytes", lambda data, context: calls.append(len(data)) or f"OCR sentence {len(calls)}.")
    text = extract(engine, source, target, "tiff", {"ocr_language": "eng"})
    assert "OCR sentence 1." in text and "OCR sentence 2." in text and len(calls) == 2
